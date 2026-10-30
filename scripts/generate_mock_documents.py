import argparse
import sys
from pathlib import Path

# Add project root to sys.path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import docx
import pymupdf as fitz


def generate_mock_markdown(output_path: Path):
    """Generates a structured markdown test document with headings, tables, and page breaks."""
    content = """# Global Information Security Standard

<!-- Page 1 -->
## 1. Access Control Policy
All employees and contractors accessing enterprise production environments must enforce hardware-backed multi-factor authentication (MFA).

### 1.1 Password & Token Requirements
Passwords must be at least 16 characters in length and rotated every 90 days. API tokens must never be hardcoded in version control repositories.

<!-- Page 2 -->
## 2. Infrastructure Security Tiers
The following matrix details replication targets, encryption standards, and backup frequencies:

| Deployment Tier | Environment | Encryption Standard | Replication Factor | RPO Target |
| --- | --- | --- | --- | --- |
| Tier-1 | Production Core | AES-256-GCM | 3x Multi-Region | < 1 minute |
| Tier-2 | Production Analytics | AES-256-CBC | 2x Single-Region | < 1 hour |
| Tier-3 | Staging / QA | Platform Default | 1x Local | < 24 hours |

### 2.1 Compliance Deviations
Any exception to this tiering standard requires formal sign-off from the Chief Information Security Officer.
"""
    output_path.write_text(content, encoding="utf-8")
    print(f"[Created MD]   {output_path.name}")


def generate_mock_text(output_path: Path):
    """Generates a plain text document with heuristic headings, ASCII tables, and form feeds."""
    content = """EMPLOYEE EXPENSE AND REIMBURSEMENT POLICY
==========================================

This document defines standard business reimbursement guidelines.

1. SUBMISSION GUIDELINES
All travel expense claims must be submitted within 30 days of trip completion.
Original digital receipts must accompany any individual item exceeding $25.00.

1.1 Per Diem Allowance
Reimbursement rates depend on the operational jurisdiction.
\f
2. DAILY PER DIEM CAPS
The following schedule outlines maximum allowable expenses per day:

| Category | Tier 1 City | Tier 2 City | Tier 3 City |
| --- | --- | --- | --- |
| Breakfast | $25.00 | $20.00 | $15.00 |
| Lunch | $35.00 | $30.00 | $25.00 |
| Dinner | $65.00 | $50.00 | $40.00 |
| Incidentals | $15.00 | $10.00 | $10.00 |

Claims exceeding these amounts will be automatically routed for senior managerial review.
"""
    output_path.write_text(content, encoding="utf-8")
    print(f"[Created TXT]  {output_path.name}")


def generate_mock_docx(output_path: Path):
    """Generates a Microsoft Word .docx file with style-based headings and a formatted table."""
    doc = docx.Document()

    # Title & Subtitle
    doc.add_heading("Cloud Architecture Blueprint", level=0)
    doc.add_paragraph("Enterprise Engineering Architecture & System Guidelines")

    # Section 1
    doc.add_heading("1. Microservice Topology", level=1)
    doc.add_paragraph(
        "All customer-facing applications are deployed as containerized microservices "
        "orchestrated through Amazon EKS across multiple availability zones."
    )

    doc.add_heading("1.1 Network Ingress", level=2)
    doc.add_paragraph(
        "Ingress traffic is secured via AWS Application Load Balancer with CloudFront CDN caching."
    )

    # Section 2 & Table
    doc.add_heading("2. Storage and Database Allocation", level=1)
    doc.add_paragraph(
        "The following database matrix governs technology selection for new persistence layers:"
    )

    # Add Table
    table = doc.add_table(rows=4, cols=4)
    table.style = "Table Grid"

    headers = ["Workload Type", "Primary Engine", "Storage Engine", "High Availability"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h

    data = [
        ["Transactional Relational", "PostgreSQL 16", "Amazon Aurora", "Multi-AZ Cluster"],
        ["Key-Value Cache", "Redis 7.2", "ElastiCache Memory", "Cluster Multi-Replica"],
        ["Vector Embeddings", "PostgreSQL + pgvector", "NVMe SSD EBS", "Read Replicas"],
    ]

    for row_idx, row_data in enumerate(data):
        for col_idx, cell_value in enumerate(row_data):
            table.rows[row_idx + 1].cells[col_idx].text = cell_value

    doc.save(str(output_path))
    print(f"[Created DOCX] {output_path.name}")


def generate_mock_pdf(output_path: Path):
    """Generates a multi-page PDF document with styled headings and vector-bordered tables."""
    doc = fitz.open()

    # Page 1: Title & Headings
    page1 = doc.new_page()
    page1.insert_text(fitz.Point(50, 70), "Annual Engineering Milestones", fontsize=22)
    page1.insert_text(fitz.Point(50, 110), "1. Executive Summary", fontsize=15)
    page1.insert_text(
        fitz.Point(50, 135),
        "During this fiscal cycle, the Document Intelligence Pipeline achieved 99.95% uptime.",
        fontsize=10,
    )
    page1.insert_text(fitz.Point(50, 170), "1.1 Scalability Improvements", fontsize=13)
    page1.insert_text(
        fitz.Point(50, 195),
        "Worker pool concurrency was optimized, reducing median parsing latency to under 200ms.",
        fontsize=10,
    )

    # Page 2: Headings & Drawn Table
    page2 = doc.new_page()
    page2.insert_text(fitz.Point(50, 70), "2. Pipeline Throughput Benchmarks", fontsize=15)
    page2.insert_text(
        fitz.Point(50, 95),
        "The following table summarizes throughput and latency across document formats:",
        fontsize=10,
    )

    # Draw table borders
    table_rect = fitz.Rect(50, 110, 480, 190)
    page2.draw_rect(table_rect)
    # Header divider
    page2.draw_line(fitz.Point(50, 130), fitz.Point(480, 130))
    # Row dividers
    page2.draw_line(fitz.Point(50, 150), fitz.Point(480, 150))
    page2.draw_line(fitz.Point(50, 170), fitz.Point(480, 170))
    # Col dividers
    page2.draw_line(fitz.Point(160, 110), fitz.Point(160, 190))
    page2.draw_line(fitz.Point(280, 110), fitz.Point(280, 190))
    page2.draw_line(fitz.Point(380, 110), fitz.Point(380, 190))

    # Text cells
    page2.insert_text(fitz.Point(60, 125), "Document Format", fontsize=9)
    page2.insert_text(fitz.Point(170, 125), "Avg Tokens", fontsize=9)
    page2.insert_text(fitz.Point(290, 125), "Parse Time", fontsize=9)
    page2.insert_text(fitz.Point(390, 125), "Chunk Count", fontsize=9)

    page2.insert_text(fitz.Point(60, 145), "PDF (PyMuPDF)", fontsize=9)
    page2.insert_text(fitz.Point(170, 145), "2,400", fontsize=9)
    page2.insert_text(fitz.Point(290, 145), "18 ms", fontsize=9)
    page2.insert_text(fitz.Point(390, 145), "4", fontsize=9)

    page2.insert_text(fitz.Point(60, 165), "DOCX (python-docx)", fontsize=9)
    page2.insert_text(fitz.Point(170, 165), "1,850", fontsize=9)
    page2.insert_text(fitz.Point(290, 165), "12 ms", fontsize=9)
    page2.insert_text(fitz.Point(390, 165), "3", fontsize=9)

    page2.insert_text(fitz.Point(60, 185), "Markdown (MD)", fontsize=9)
    page2.insert_text(fitz.Point(170, 185), "1,200", fontsize=9)
    page2.insert_text(fitz.Point(290, 185), "4 ms", fontsize=9)
    page2.insert_text(fitz.Point(390, 185), "2", fontsize=9)

    doc.save(str(output_path))
    doc.close()
    print(f"[Created PDF]  {output_path.name}")


def main():
    parser = argparse.ArgumentParser(description="Generate mock documents for local pipeline testing.")
    parser.add_argument(
        "--output-dir",
        type=str,
        default="data/mock_docs",
        help="Target directory to write mock documents",
    )
    args = parser.parse_args()

    out_dir = Path(args.output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    print("=" * 70)
    print(f"Generating Mock Documents into: {out_dir.resolve()}")
    print("=" * 70)

    generate_mock_markdown(out_dir / "sample_policy.md")
    generate_mock_text(out_dir / "sample_policy.txt")
    generate_mock_docx(out_dir / "sample_architecture.docx")
    generate_mock_pdf(out_dir / "sample_milestones.pdf")

    print("=" * 70)
    print("Mock document generation complete! Ready for local ingestion testing.")
    print("=" * 70)


if __name__ == "__main__":
    main()
