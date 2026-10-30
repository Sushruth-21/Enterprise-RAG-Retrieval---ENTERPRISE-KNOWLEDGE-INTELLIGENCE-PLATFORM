import io
from pathlib import Path
import docx
import pymupdf as fitz
import pytest

from parsers import (
    DocxParser,
    ElementType,
    MarkdownParser,
    PDFParser,
    ParsedDocument,
    TextParser,
    get_parser,
    parse_document,
)


@pytest.fixture
def sample_dir(tmp_path: Path) -> Path:
    return tmp_path


# =========================================================================
# 1. Markdown Parser Tests
# =========================================================================
def test_markdown_parser_headings_and_tables(sample_dir: Path):
    md_content = """# Company Security Policy

This document defines core information security guidelines.

<!-- Page 1 -->
## 1. Access Control
Access to production systems requires multi-factor authentication.

### 1.1 Password Requirements
Passwords must be at least 16 characters long.

<!-- Page 2 -->
## 2. Infrastructure Inventory
Below is the infrastructure tier allocation table:

| Service | Environment | Tier | Instances |
| --- | --- | --- | --- |
| auth-service | prod | Tier-1 | 4 |
| payment-api | prod | Tier-1 | 6 |
| analytics-worker | staging | Tier-3 | 2 |

Final paragraph concluding policy.
"""
    file_path = sample_dir / "security_policy.md"
    file_path.write_text(md_content, encoding="utf-8")

    parser = MarkdownParser()
    doc = parser.parse(file_path)

    assert doc.success is True
    assert doc.title == "Company Security Policy"
    assert doc.total_pages == 2

    # Check headings
    headings = doc.get_headings()
    assert len(headings) >= 3
    assert headings[0].content == "Company Security Policy"
    assert headings[0].element_type == ElementType.HEADING_1
    assert headings[1].content == "1. Access Control"
    assert headings[1].element_type == ElementType.HEADING_2
    assert headings[1].heading_hierarchy == ["Company Security Policy", "1. Access Control"]
    assert headings[2].content == "1.1 Password Requirements"
    assert headings[2].element_type == ElementType.HEADING_3
    assert headings[2].heading_hierarchy == [
        "Company Security Policy",
        "1. Access Control",
        "1.1 Password Requirements",
    ]

    # Check tables
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert table.page_number == 2
    assert table.headers == ["Service", "Environment", "Tier", "Instances"]
    assert len(table.rows) == 3
    assert "| auth-service | prod | Tier-1 | 4 |" in table.markdown
    assert "| --- | --- | --- | --- |" in table.markdown


# =========================================================================
# 2. DOCX Parser Tests
# =========================================================================
def test_docx_parser_headings_and_tables(sample_dir: Path):
    file_path = sample_dir / "sample_doc.docx"
    doc_builder = docx.Document()

    doc_builder.add_heading("Corporate Travel Policy", level=0)
    p1 = doc_builder.add_paragraph("All employees must follow standard reimbursement guidelines.")

    doc_builder.add_heading("1. Per Diem Rates", level=1)
    p2 = doc_builder.add_paragraph("Rates depend on local jurisdiction.")

    doc_builder.add_heading("1.1 Allowable Expenses", level=2)
    # Add Table
    table = doc_builder.add_table(rows=3, cols=2)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Category"
    hdr_cells[1].text = "Daily Limit"

    row1 = table.rows[1].cells
    row1[0].text = "Meals"
    row1[1].text = "$75"

    row2 = table.rows[2].cells
    row2[0].text = "Lodging"
    row2[1].text = "$250"

    doc_builder.save(str(file_path))

    parser = DocxParser()
    parsed_doc = parser.parse(file_path)

    assert parsed_doc.success is True
    assert "Corporate Travel Policy" in (parsed_doc.title or "")

    # Headings
    headings = parsed_doc.get_headings()
    assert any("Per Diem Rates" in h.content for h in headings)
    assert any("Allowable Expenses" in h.content for h in headings)

    # Tables
    assert len(parsed_doc.tables) == 1
    t = parsed_doc.tables[0]
    assert t.headers == ["Category", "Daily Limit"]
    assert len(t.rows) == 2
    assert "| Category | Daily Limit |" in t.markdown
    assert "| Meals | $75 |" in t.markdown


# =========================================================================
# 3. Plain Text Parser Tests
# =========================================================================
def test_text_parser_structure(sample_dir: Path):
    txt_content = """EMPLOYEE ONBOARDING GUIDE
=========================

Welcome to the company.

1. FIRST WEEK CHECKLIST
Complete IT security setup and HR orientation.

1.1 Credentials
Receive corporate hardware and SSO tokens.
\f
2. HARDWARE INVENTORY
| Item | Model | Quantity |
| --- | --- | --- |
| Laptop | ThinkPad X1 | 1 |
| Monitor | Dell 27 | 2 |
"""
    file_path = sample_dir / "onboarding.txt"
    file_path.write_text(txt_content, encoding="utf-8")

    parser = TextParser()
    parsed_doc = parser.parse(file_path)

    assert parsed_doc.success is True
    assert parsed_doc.total_pages == 2
    assert "EMPLOYEE ONBOARDING GUIDE" in (parsed_doc.title or "")

    headings = parsed_doc.get_headings()
    assert len(headings) >= 2
    assert any("FIRST WEEK CHECKLIST" in h.content for h in headings)

    # Check table
    assert len(parsed_doc.tables) == 1
    assert parsed_doc.tables[0].page_number == 2
    assert parsed_doc.tables[0].headers == ["Item", "Model", "Quantity"]


# =========================================================================
# 4. PyMuPDF PDF Parser Tests (Font-size hierarchy & page preservation)
# =========================================================================
def test_pdf_parser_with_pymupdf(sample_dir: Path):
    pdf_path = sample_dir / "test_document.pdf"

    # Create a synthetic 2-page PDF with PyMuPDF
    doc = fitz.open()

    # Page 1: Title and Headings
    page1 = doc.new_page()
    page1.insert_text(fitz.Point(50, 60), "Annual Technical Report", fontsize=24)
    page1.insert_text(fitz.Point(50, 100), "1. Executive Summary", fontsize=16)
    page1.insert_text(
        fitz.Point(50, 130),
        "This report outlines our engineering milestones and pipeline architecture.",
        fontsize=10,
    )
    page1.insert_text(fitz.Point(50, 170), "1.1 Core Systems", fontsize=13)
    page1.insert_text(
        fitz.Point(50, 195),
        "All ingestion workloads are distributed across Celery workers.",
        fontsize=10,
    )

    # Page 2: Headings and Table drawn with vector lines
    page2 = doc.new_page()
    page2.insert_text(fitz.Point(50, 60), "2. Performance Metrics", fontsize=16)

    # Draw table borders
    # Outer box
    page2.draw_rect(fitz.Rect(50, 80, 450, 140))
    # Horizontal line between header and data
    page2.draw_line(fitz.Point(50, 100), fitz.Point(450, 100))
    page2.draw_line(fitz.Point(50, 120), fitz.Point(450, 120))
    # Vertical line separating cols
    page2.draw_line(fitz.Point(250, 80), fitz.Point(250, 140))

    # Cell texts
    page2.insert_text(fitz.Point(60, 95), "Metric", fontsize=9)
    page2.insert_text(fitz.Point(260, 95), "Value", fontsize=9)
    page2.insert_text(fitz.Point(60, 115), "Throughput", fontsize=9)
    page2.insert_text(fitz.Point(260, 115), "500 docs/min", fontsize=9)
    page2.insert_text(fitz.Point(60, 135), "P99 Latency", fontsize=9)
    page2.insert_text(fitz.Point(260, 135), "180 ms", fontsize=9)

    doc.save(str(pdf_path))
    doc.close()

    parser = PDFParser(preferred_engine="pymupdf")
    parsed_doc = parser.parse(pdf_path)

    assert parsed_doc.success is True
    assert parsed_doc.total_pages == 2
    assert "Annual Technical Report" in parsed_doc.raw_text

    # Verify headings and page mapping
    headings = parsed_doc.get_headings()
    assert len(headings) >= 2
    page1_elements = parsed_doc.get_elements_by_page(1)
    page2_elements = parsed_doc.get_elements_by_page(2)
    assert len(page1_elements) > 0
    assert len(page2_elements) > 0

    # Verify tables
    assert len(parsed_doc.tables) >= 1
    t = parsed_doc.tables[0]
    assert t.page_number == 2
    assert "Metric" in t.markdown or "Throughput" in t.markdown
    assert "---" in t.markdown


# =========================================================================
# 5. Error Handling for Malformed Documents
# =========================================================================
def test_error_handling_malformed_documents(sample_dir: Path):
    # Corrupted PDF (random garbage bytes)
    corrupted_pdf = sample_dir / "corrupted.pdf"
    corrupted_pdf.write_bytes(b"%PDF-1.4\nGARBAGE_BYTES_CORRUPTED_STREAM\x00\xff\xee")

    parser = PDFParser()
    result = parser.parse(corrupted_pdf)

    # Must log failure and return ParsedDocument with success=False WITHOUT crashing
    assert isinstance(result, ParsedDocument)
    assert result.success is False
    assert result.error_message is not None
    assert "Extraction failure" in result.error_message or "corrupted" in result.error_message.lower()

    # Non-existent file
    non_existent = sample_dir / "does_not_exist.pdf"
    res_missing = parser.parse(non_existent)
    assert res_missing.success is False
    assert "not found" in res_missing.error_message.lower()


# =========================================================================
# 6. Factory Routing Tests
# =========================================================================
def test_parser_factory_routing(sample_dir: Path):
    f_pdf = sample_dir / "doc.pdf"
    f_docx = sample_dir / "doc.docx"
    f_txt = sample_dir / "doc.txt"
    f_md = sample_dir / "doc.md"

    assert isinstance(get_parser(f_pdf), PDFParser)
    assert isinstance(get_parser(f_docx), DocxParser)
    assert isinstance(get_parser(f_txt), TextParser)
    assert isinstance(get_parser(f_md), MarkdownParser)
